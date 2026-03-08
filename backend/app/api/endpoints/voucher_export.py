from io import BytesIO
from typing import List
from fastapi import APIRouter, Depends, Response
from sqlalchemy.orm import Session
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

from app.core.database import get_db
from app.core.permissions import require_permission
from app.models.voucher import Voucher

router = APIRouter(prefix="/vouchers", tags=["凭证导出"])


# 尝试注册中文字体
try:
    pdfmetrics.registerFont(TTFont('SimSun', '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'))
    CHINESE_FONT = 'SimSun'
except:
    try:
        pdfmetrics.registerFont(TTFont('SimSun', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
        CHINESE_FONT = 'SimSun'
    except:
        CHINESE_FONT = 'Helvetica'


def generate_voucher_pdf(voucher: Voucher) -> bytes:
    """生成凭证 PDF"""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
    
    elements = []
    styles = getSampleStyleSheet()
    
    # 标题
    title_style = styles['Title']
    title_style.fontName = CHINESE_FONT
    title_style.fontSize = 18
    elements.append(Paragraph("记账凭证", title_style))
    elements.append(Spacer(1, 20))
    
    # 表头信息
    header_style = styles['Normal']
    header_style.fontName = CHINESE_FONT
    header_style.fontSize = 10
    
    header_text = f"日期: {voucher.voucher_date}    凭证号: {voucher.voucher_no}    会计期间: {voucher.period}"
    elements.append(Paragraph(header_text, header_style))
    elements.append(Spacer(1, 10))
    
    # 分录表格
    table_data = [['摘要', '会计科目', '借方金额', '贷方金额']]
    
    total_debit = 0
    total_credit = 0
    
    for item in voucher.items:
        debit = float(item.debit_amount) if item.debit_amount else 0
        credit = float(item.credit_amount) if item.credit_amount else 0
        total_debit += debit
        total_credit += credit
        
        table_data.append([
            item.summary or '',
            f"{item.subject_code} {item.subject_name}",
            f"{debit:,.2f}" if debit > 0 else '',
            f"{credit:,.2f}" if credit > 0 else ''
        ])
    
    # 合计行
    table_data.append(['合计', '', f"{total_debit:,.2f}", f"{total_credit:,.2f}"])
    
    # 创建表格
    table = Table(table_data, colWidths=[150, 180, 80, 80])
    table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), CHINESE_FONT),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('ALIGN', (0, 0), (0, -1), 'LEFT'),
        ('ALIGN', (1, 0), (1, -1), 'LEFT'),
        ('ALIGN', (2, 0), (-1, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('BACKGROUND', (0, -1), (-1, -1), colors.lightgrey),
        ('FONTSIZE', (0, -1), (-1, -1), 10),
    ]))
    
    elements.append(table)
    elements.append(Spacer(1, 20))
    
    # 摘要
    if voucher.summary:
        elements.append(Paragraph(f"摘要: {voucher.summary}", header_style))
        elements.append(Spacer(1, 20))
    
    # 签字栏
    footer_text = "主管: ____________    记账: ____________    出纳: ____________    审核: ____________"
    elements.append(Paragraph(footer_text, header_style))
    
    doc.build(elements)
    
    return buffer.getvalue()


def generate_voucher_excel(voucher: Voucher) -> bytes:
    """生成凭证 Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = voucher.voucher_no or "凭证"
    
    # 标题
    ws.merge_cells('A1:D1')
    ws['A1'] = '记账凭证'
    ws['A1'].font = Font(size=16, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # 表头信息
    ws['A2'] = f'日期: {voucher.voucher_date}'
    ws['B2'] = f'凭证号: {voucher.voucher_no}'
    ws['C2'] = f'会计期间: {voucher.period}'
    
    # 摘要
    if voucher.summary:
        ws['A3'] = f'摘要: {voucher.summary}'
    
    # 表头
    headers = ['摘要', '会计科目', '借方金额', '贷方金额']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    # 数据行
    row = 6
    total_debit = 0
    total_credit = 0
    
    for item in voucher.items:
        ws.cell(row=row, column=1, value=item.summary)
        ws.cell(row=row, column=2, value=f"{item.subject_code} {item.subject_name}")
        
        debit = float(item.debit_amount) if item.debit_amount else 0
        credit = float(item.credit_amount) if item.credit_amount else 0
        total_debit += debit
        total_credit += credit
        
        if debit > 0:
            ws.cell(row=row, column=3, value=debit)
        if credit > 0:
            ws.cell(row=row, column=4, value=credit)
        
        row += 1
    
    # 合计行
    ws.cell(row=row, column=1, value='合计')
    ws.cell(row=row, column=3, value=total_debit)
    ws.cell(row=row, column=4, value=total_credit)
    
    # 设置边框
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for r in ws.iter_rows(min_row=5, max_row=row, min_col=1, max_col=4):
        for cell in r:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
    
    # 调整列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return buffer.read()


@router.get("/{voucher_id}/export")
async def export_voucher(
    voucher_id: str,
    format: str = "pdf",
    current_user=Depends(require_permission("vouchers:read")),
    db: Session = Depends(get_db)
):
    """导出单个凭证"""
    voucher = db.query(Voucher).filter(Voucher.id == voucher_id).first()
    if not voucher:
        from app.core.response import error_response
        return error_response(404, "凭证不存在")
    
    if str(voucher.customer_id) != str(current_user.customer_id) and current_user.role != "admin":
        from app.core.response import error_response
        return error_response(403, "无权访问")
    
    if format == "pdf":
        content = generate_voucher_pdf(voucher)
        filename = f"{voucher.voucher_no or voucher_id}.pdf"
        media_type = "application/pdf"
    elif format == "excel":
        content = generate_voucher_excel(voucher)
        filename = f"{voucher.voucher_no or voucher_id}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        from app.core.response import error_response
        return error_response(400, "不支持的格式")
    
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
